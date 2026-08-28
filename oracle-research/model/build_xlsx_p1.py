import openpyxl, json, datetime
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL
import model as M

BLUE=Font(name='Arial',size=10,color='0000FF')          # hardcoded input / scenario lever
BLACK=Font(name='Arial',size=10)                          # formula
GREEN=Font(name='Arial',size=10,color='008000')           # link to another sheet
HDR=Font(name='Arial',size=10,bold=True)
TITLE=Font(name='Arial',size=13,bold=True)
SUB=Font(name='Arial',size=10,bold=True,color='444444')
NOTE=Font(name='Arial',size=8,italic=True,color='666666')
YFILL=PatternFill('solid',fgColor='FFFF00')
GFILL=PatternFill('solid',fgColor='E8E8E8')
HFILL=PatternFill('solid',fgColor='1F3864')
WHDR=Font(name='Arial',size=10,bold=True,color='FFFFFF')
THIN=Border(bottom=Side(style='thin',color='999999'))
CUR='$#,##0;($#,##0);-'
CUR2='$#,##0.00;($#,##0.00);-'
PCT='0.0%;(0.0%);-'
PCT2='0.00%;(0.00%);-'
MULT='0.0x'
NUM='#,##0;(#,##0);-'

wb=openpyxl.Workbook()
FY=[2027,2028,2029,2030,2031]

def sheet(name):
    ws=wb.create_sheet(name); ws.sheet_view.showGridLines=False
    return ws
def title(ws,t,sub=None):
    ws['A1']=t; ws['A1'].font=TITLE
    if sub: ws['A2']=sub; ws['A2'].font=NOTE
def hrow(ws,r,labels,start=1,fill=True):
    for i,l in enumerate(labels):
        c=ws.cell(row=r,column=start+i,value=l); c.font=WHDR if fill else HDR
        if fill: c.fill=HFILL
        c.alignment=Alignment(horizontal='center' if i else 'left')

# ============================================================== README
ws=wb.active; ws.title='README'; ws.sheet_view.showGridLines=False
title(ws,'Oracle Corporation (ORCL) — independent forecast & valuation model')
rows=[
 ('',''),
 ('Prepared for','Joshua Lacroix'),
 ('Prepared by','Independent build (Claude Code). Not investment advice.'),
 ('Model date','2026-08-28'),
 ('Market data as of','2026-08-27 close. ORCL $151.94.'),
 ('',''),
 ('PRIMARY SOURCES',''),
 ('FY2026 Form 10-K','SEC EDGAR CIK 0001341439, accession 0001193125-26-277521, filed 2026-06-22, FYE 2026-05-31'),
 ('Q4/FY26 press release','Exhibit 99.1 to Form 8-K filed 2026-06-10, accession 0001193125-26-265848'),
 ('Market / comps data','stockanalysis.com, pulled 2026-08-27/28 (single source, single date)'),
 ('Risk-free rate','10-year US Treasury (^TNX) 4.672%, 2026-08-27'),
 ('',''),
 ('WHAT IS *NOT* SOURCED HERE',''),
 ('Q4 FY26 slide deck / call transcript','investor.oracle.com returned HTTP 403. Any figure originating only from the call '
   '(FY27 net capex ~$70bn, "no additional debt in calendar 2026", component cost pass-through) is '
   'flagged UNVERIFIED SECONDARY and is NOT used as a model input.'),
 ('FY2028+ consensus','Paywalled at source. FY28 estimates in this model are the analyst\'s own, not consensus.'),
 ('',''),
 ('COLOUR LEGEND',''),
 ('Blue text','Hardcoded input / scenario lever — edit these'),
 ('Black text','Formula calculated on this sheet'),
 ('Green text','Link to another sheet in this workbook'),
 ('Yellow fill','Key assumption. The answer is most sensitive to these.'),
 ('',''),
 ('HOW TO DRIVE THE MODEL',''),
 ('1.','Set the scenario on Inputs!B4 to Base, Bull or Bear. Everything downstream re-solves.'),
 ('2.','Edit any blue cell in the BASE / BULL / BEAR blocks on the Inputs tab. Never edit a black cell.'),
 ('3.','Every driver row carries a source note in column J of the Inputs tab.'),
 ('',''),
 ('TAB MAP',''),
 ('Inputs','All drivers, three scenarios, with source notes'),
 ('Historicals','FY2024–FY2026 actuals, every line tied to the 10-K'),
 ('RPO_Schedule','The disclosed 12% / 34% / 34% / 20% conversion schedule and the revenue build off it'),
 ('Capex_Depreciation','Capex→CIP→in-service→depreciation vintage waterfall. The most important mechanical tab.'),
 ('Income_Statement','FY24A–FY31E, GAAP'),
 ('NonGAAP_Recon','Full GAAP↔non-GAAP bridge including the tax-rate wedge'),
 ('Cash_Flow','OCF, capex (three definitions), FCF'),
 ('Debt_Shares','Debt schedule, interest, share count, ATM, mandatory convertible'),
 ('Valuation_RPO_NPV','PRIMARY valuation: contract-level NPV of the $638bn backlog'),
 ('Valuation_DCF','Cross-check: unlevered FCF DCF and FCFE'),
 ('Sensitivity','Two-way tables and driver ranking'),
 ('ROIC_Breakeven','Required OCI gross margin to clear WACC. The central question.'),
 ('Comps_Software','Enterprise software comp set, calendarised'),
 ('Comps_AI','AI infrastructure comp set, calendarised'),
 ('Formula_Audit','Every simplification, circularity break and known limitation'),
]
r=4
for a,b in rows:
    ws.cell(row=r,column=1,value=a).font=HDR if (b and not a.startswith(('1.','2.','3.'))) and a.isupper() else BLACK
    if a and a.isupper(): ws.cell(row=r,column=1).font=SUB
    c=ws.cell(row=r,column=2,value=b); c.font=BLACK; c.alignment=Alignment(wrap_text=True,vertical='top')
    r+=1
ws.column_dimensions['A'].width=30; ws.column_dimensions['B'].width=110

# ============================================================== INPUTS
ws=sheet('Inputs')
title(ws,'INPUTS — every driver, three scenarios','Blue = editable. Column J carries the source or the basis for the estimate. E = analyst estimate, D = disclosed by Oracle.')
ws['A4']='SCENARIO'; ws['A4'].font=HDR
ws['B4']='Base'; ws['B4'].font=BLUE; ws['B4'].fill=YFILL
ws['C4']='<- type Base, Bull or Bear'; ws['C4'].font=NOTE
ws['A5']='scenario index'; ws['A5'].font=NOTE
ws['B5']='=MATCH($B$4,$A$996:$A$998,0)'; ws['B5'].font=BLACK
ws['A996']='Base'; ws['A997']='Bull'; ws['A998']='Bear'
for rr in (996,997,998): ws.cell(row=rr,column=1).font=NOTE

# driver spec: (key,label,single?,fmt,base,bull,bear,source)
D=[]
def add(key,label,single,fmt,b,u,e,src): D.append((key,label,single,fmt,b,u,e,src))
S=M.BASE; U=M.BULL; E=M.BEAR
add('SEC_RPO','— RPO CONVERSION —',None,None,None,None,None,'')
add('rpo26','RPO at 5/31/2026 ($m)',True,NUM,638000,638000,638000,'D: 10-K Note 1 & MD&A. $638bn vs $137.8bn at 5/31/25.')
add('sch_y1','Disclosed: recognised over next 12 months',True,PCT,0.12,0.12,0.12,'D: 10-K Note 1 verbatim.')
add('sch_y23','Disclosed: months 13–36',True,PCT,0.34,0.34,0.34,'D: 10-K Note 1 verbatim.')
add('sch_y45','Disclosed: months 37–60',True,PCT,0.34,0.34,0.34,'D: 10-K Note 1 verbatim.')
add('sch_tail','Disclosed: thereafter (residual)',True,PCT,0.20,0.20,0.20,'D: "and the remainder thereafter".')
add('rpo_legacy','"Legacy" (pre-AI) backlog ($m)',True,NUM,S['rpo_legacy'],U['rpo_legacy'],E['rpo_legacy'],'E: ≈ the 5/31/25 RPO of $137.8bn. Splits the backlog into fast-converting legacy and slow-converting AI.')
add('rpo_haircut','AI backlog never recognised (counterparty)',True,PCT,S['rpo_haircut'],U['rpo_haircut'],E['rpo_haircut'],'E: KEY LEVER. 10-K discloses NO customer >10% of revenue and does not quantify RPO concentration.')
add('legacy_conv','Legacy backlog conversion, % per year',False,PCT,S['legacy_conv'],U['legacy_conv'],E['legacy_conv'],'E: pre-AI RPO converted ~60–62% in 12 months when last tagged (FY2021 10-K).')
add('ai_y2_share','Share of the m13–36 bucket falling in FY28',True,PCT,S['ai_split_y23'][0],U['ai_split_y23'][0],E['ai_split_y23'][0],'E: intra-bucket timing. Oracle discloses the 24-month total only.')
add('ai_y4_share','Share of the m37–60 bucket falling in FY30',True,PCT,S['ai_split_y45'][0],U['ai_split_y45'][0],E['ai_split_y45'][0],'E: intra-bucket timing.')
add('new_bookings','New bookings signed after 5/31/26 ($m)',False,NUM,S['new_bookings'],U['new_bookings'],E['new_bookings'],'E: PURE ESTIMATE. Dominates FY30–31. Oracle added ~$500bn of RPO in FY26 alone; not repeatable.')
add('new_conv','New-bookings conversion, % per year',False,PCT,S['new_conv'],U['new_conv'],E['new_conv'],'E: assumes new contracts convert on a similar back-loaded profile to the AI backlog.')
add('nonrpo_base','Revenue not covered by RPO, FY27 ($m)',True,NUM,S['nonrpo_base'],U['nonrpo_base'],E['nonrpo_base'],'E: hardware, non-committed consumption, some services/new licence.')
add('nonrpo_g','  growth',True,PCT,S['nonrpo_growth'],U['nonrpo_growth'],E['nonrpo_growth'],'E.')
add('SEC_SEG','— SEGMENT DRIVERS (IaaS is the residual) —',None,None,None,None,None,'')
add('saas_g','Cloud Applications (SaaS) growth',False,PCT,S['saas_g'],U['saas_g'],E['saas_g'],'E: FY26A +11.3%. Mgmt guides Oracle Health to double-digit growth in FY27 (PR 6/10/26).')
add('lic_g','Software licence growth',False,PCT,S['lic_g'],U['lic_g'],E['lic_g'],'E: FY24–26A $5,081→$5,201→$4,737m, structurally declining.')
add('sup_g','Software support growth',False,PCT,S['sup_g'],U['sup_g'],E['sup_g'],'E: FY24–26A $19,609→$19,523→$19,804m. A genuine annuity.')
add('svc_g','Services growth',False,PCT,S['svc_g'],U['svc_g'],E['svc_g'],'E: FY26A +9.7%.')
add('hw_g','Hardware growth',False,PCT,S['hw_g'],U['hw_g'],E['hw_g'],'E: FY26A +5.0%.')
add('SEC_MGN','— MARGINS —',None,None,None,None,None,'')
add('iaas_ndc','IaaS cost ex-depreciation, % of IaaS revenue',False,PCT,S['iaas_nondep_cost'],U['iaas_nondep_cost'],E['iaas_nondep_cost'],'E: KEY LEVER. Power, bandwidth, site labour, component pass-through. Oracle does not disclose OCI cost structure.')
add('oci_dep_sh','Share of total depreciation charged to IaaS COGS',False,PCT,S['ocidep_share'],U['ocidep_share'],E['ocidep_share'],'E: not disclosed. Depreciation is inside cost of revenue as Oracle reports it.')
add('saas_gm','SaaS gross margin',False,PCT,S['saas_gm'],U['saas_gm'],E['saas_gm'],'E: not disclosed separately.')
add('sup_gm','Software support gross margin',False,PCT,S['sup_gm'],U['sup_gm'],E['sup_gm'],'E: not disclosed separately.')
add('lic_gm','Software licence gross margin',False,PCT,S['lic_gm'],U['lic_gm'],E['lic_gm'],'E: not disclosed separately.')
add('svc_gm','Services gross margin',False,PCT,S['svc_gm'],U['svc_gm'],E['svc_gm'],'D: FY26A 20.7% ($5,743m rev / $4,556m cost, 10-K R4).')
add('hw_gm','Hardware gross margin',False,PCT,S['hw_gm'],U['hw_gm'],E['hw_gm'],'D: FY26A 71.9% ($3,084m / $868m, 10-K R4).')
add('SEC_OPX','— OPERATING EXPENSE —',None,None,None,None,None,'')
add('sm_pct','Sales & marketing, % of revenue',False,PCT,S['sm_pct'],U['sm_pct'],E['sm_pct'],'E: FY26A 12.4%. Scale leverage as revenue steps up.')
add('rd_g','R&D growth',False,PCT,S['rd_growth'],U['rd_growth'],E['rd_growth'],'E: FY26A $10,272m, +4.2%. 43,000 of 141,000 employees are in R&D (10-K Item 1).')
add('ga_g','G&A growth',False,PCT,S['ga_growth'],U['ga_growth'],E['ga_growth'],'E.')
add('sbc_g','Stock-based compensation growth',False,PCT,S['sbc_growth'],U['sbc_growth'],E['sbc_growth'],'E: FY26A $4,811m. 10-K MD&A: annualised net dilution 1.0%/yr since 6/1/2023.')
add('restr','Restructuring & other ($m)',False,NUM,S['restructuring'],U['restructuring'],E['restructuring'],'E: FY26A $1,838m (headcount 154k→141k). 10-K R59.')
add('impair','Asset impairment ($m) — bear only',False,NUM,S['impairment'],U['impairment'],E['impairment'],'E: bear case only. Written-down assets stop depreciating (relieved over 5 yrs).')
add('amort','Intangible amortisation ($m)',False,NUM,[731,694,620,582,377],[731,694,620,582,377],[731,694,620,582,377],'D: 10-K R53 — the full forward schedule is disclosed. No estimate required.')
add('SEC_CAP','— CAPEX / PP&E (LINKED, not a % of revenue) —',None,None,None,None,None,'')
add('turn_full','IaaS revenue per $ gross PP&E at full utilisation',True,'0.00',S['asset_turn_full'],U['asset_turn_full'],E['asset_turn_full'],'E: KEY LEVER. Not disclosed. FY26A implies ~0.27 on in-service PP&E, but much of that base is legacy/non-OCI.')
add('util','Utilisation of installed capacity',False,PCT,S['utilization'],U['utilization'],E['utilization'],'E: KEY LEVER. Not disclosed anywhere.')
add('cust_fund','Capacity funded by customers (prepay + customer GPUs)',False,PCT,S['cust_funded_capacity'],U['cust_funded_capacity'],E['cust_funded_capacity'],'Anchored to D: "prepaid and customer supplied hardware portions ... now total $75 billion" (PR 6/10/26).')
add('maint_pct','Maintenance capex, % of depreciation',False,PCT,S['maint_capex_pct_dep'],U['maint_capex_pct_dep'],E['maint_capex_pct_dep'],'E: rises as the fleet ages.')
add('capex_min','Committed capex floor ($m)',False,NUM,S['capex_min'],U['capex_min'],E['capex_min'],'E: GPU/server orders and data-centre leases placed 12–18 months ahead cannot be cancelled inside the year.')
add('iaas_gterm','IaaS growth assumed in FY32 (to size FY31 capex)',True,PCT,S['iaas_g_terminal'],U['iaas_g_terminal'],E['iaas_g_terminal'],'E.')
add('mix_eq','Capex mix: equipment',True,PCT,S['capex_equip'],U['capex_equip'],E['capex_equip'],'E: 10-K R49 gross PP&E is 48.6% computer/network, 17.3% buildings, 32.6% construction in progress.')
add('mix_bl','Capex mix: buildings/shell',True,PCT,S['capex_bldg'],U['capex_bldg'],E['capex_bldg'],'E.')
add('mix_land','Capex mix: land & non-depreciating',True,PCT,S['capex_land'],U['capex_land'],E['capex_land'],'E.')
add('life_eq','Equipment life (years)',True,'0.0',S['life_equip'],U['life_equip'],E['life_equip'],'D: 10-K R50 — servers and networking equipment, 6 years.')
add('life_bl','Buildings life (years)',True,'0.0',S['life_bldg'],U['life_bldg'],E['life_bldg'],'D: 10-K R49 — buildings 1–40 years. 25 used as the blended estimate.')
add('cip_open','Construction in progress at 5/31/26 ($m)',True,NUM,39973,39973,39973,'D: 10-K R49. 32.6% of gross PP&E is NOT YET DEPRECIATING.')
add('cip_rel1','CIP placed in service in FY27',True,PCT,S['cip_release'][0],U['cip_release'][0],E['cip_release'][0],'E.')
add('cip_rel2','CIP placed in service in FY28',True,PCT,S['cip_release'][1],U['cip_release'][1],E['cip_release'][1],'E.')
add('cx_isv1','Capex placed in service in year spent',True,PCT,S['capex_inservice_split'][0],U['capex_inservice_split'][0],E['capex_inservice_split'][0],'E: ~9 months average time in CIP.')
add('cx_isv2','Capex placed in service in following year',True,PCT,S['capex_inservice_split'][1],U['capex_inservice_split'][1],E['capex_inservice_split'][1],'E.')
add('base_life','Remaining life of existing in-service base (yrs)',True,'0.0',5.0,5.0,5.0,'E: FY26A depreciation $7,623m on avg in-service gross PP&E implies ~8yr blended; 5yr used on residual NBV.')
add('SEC_FIN','— FINANCING, TAX, SHARES —',None,None,None,None,None,'')
add('kd_new','Pre-tax cost of NEW debt',True,PCT2,S['cost_new_debt'],U['cost_new_debt'],E['cost_new_debt'],'E: Oracle\'s Feb-2026 30yr priced 6.70%, 40yr 6.85% (10-K MD&A). S&P cut to BBB- 7/9/26 (unverified secondary).')
add('kd_exist','Blended rate on existing debt',True,PCT2,S['rate_existing'],U['rate_existing'],E['rate_existing'],'D-derived: FY26A interest expense $4,599m on average borrowings.')
add('ki','Interest income rate on cash',True,PCT2,S['interest_income_rate'],U['interest_income_rate'],E['interest_income_rate'],'E: FY26A interest income $780m.')
add('eq_raise','Equity issuance ($m)',False,NUM,S['equity_raise'],U['equity_raise'],E['equity_raise'],'D: $20bn ATM programme entered 2/2/2026, ZERO drawn at 5/31/26 (10-K R70). Bear draws more, lower.')
add('eq_price','Assumed issue price ($)',False,CUR2,S['equity_price'],U['equity_price'],E['equity_price'],'E.')
add('tax_gaap','GAAP effective tax rate',False,PCT,S['tax_rate_gaap'],U['tax_rate_gaap'],E['tax_rate_gaap'],'E: FY26A was 12.6%, flattered by a $2,062m SBC windfall benefit (10-K R80) that will not repeat at $152/share.')
add('tax_ng','Non-GAAP tax rate',False,PCT,S['tax_rate_ng'],U['tax_rate_ng'],E['tax_rate_ng'],'D: FY26A 19.9% (press release footnote 5).')
add('div_ps','Common dividend per share ($)',False,CUR2,S['dividend_ps'],U['dividend_ps'],E['dividend_ps'],'D: $0.50/qtr declared 6/10/26, flat vs FY26. Dividend growth has stopped.')
add('dil_sbc','Annual net dilution from stock awards',True,PCT,S['dilution_sbc'],U['dilution_sbc'],E['dilution_sbc'],'D: 10-K MD&A — 1.0% annualised since 6/1/2023; max potential dilution 3.5%.')
add('prefsh','Shares on mandatory conversion (m)',True,'0.0',M.PREF['shares_if_conv_at_mkt'],M.PREF['shares_if_conv_at_mkt'],M.PREF['shares_if_conv_at_mkt'],'D: 10-K Note 10 — converts 1/15/2029 at 499.8126–624.7657 sh per pref. At $151.94 (< $160.06) the MAX ratio applies.')
add('prefdiv','Preferred dividend, annual ($m)',True,NUM,M.PREF['div_annual'],M.PREF['div_annual'],M.PREF['div_annual'],'D: 6.50% on $5.0bn = $325m. FY26A actual $103m (part-year).')
add('cust_pre','Customer prepayments in OCF ($m)',False,NUM,S['cust_prepay'],U['cust_prepay'],E['cust_prepay'],'D-anchored: FY26A $4,592m, a brand-new cash-flow line (10-K R8).')
add('st_fin','Short-term financing re capex ($m)',False,NUM,S['st_fin_capex'],U['st_fin_capex'],E['st_fin_capex'],'D-anchored: FY26A $3,345m (10-K R8, financing activities).')
add('nwc_pct','Incremental NWC, % of revenue change',True,PCT,S['nwc_pct_rev_change'],U['nwc_pct_rev_change'],E['nwc_pct_rev_change'],'E.')
add('debt_open','Total debt at 5/31/26 ($m)',True,NUM,129541,129541,129541,'D: 10-K R2 — $7,199m current + $122,342m non-current.')
add('cash_open','Cash + marketable securities at 5/31/26 ($m)',True,NUM,31894,31894,31894,'D: 10-K R2 — $31,289m + $605m.')
add('debt_mat','Scheduled debt maturities ($m)',False,NUM,[7210,10145,5500,7250,9750],[7210,10145,5500,7250,9750],[7210,10145,5500,7250,9750],'D: 10-K R57 — the full maturity schedule is disclosed.')
add('SEC_VAL','— VALUATION —',None,None,None,None,None,'')
add('wacc','WACC',True,PCT2,0.1030,0.0950,0.1109,'Derived on the Valuation_DCF tab from rf 4.672%, ERP 5.0%, beta, and a 6.75% marginal cost of debt.')
add('tg','DCF terminal growth',True,PCT,S['terminal_growth'],U['terminal_growth'],E['terminal_growth'],'E.')

hrow(ws,7,['Driver','ACTIVE (single)','FY2027','FY2028','FY2029','FY2030','FY2031','','','Source / basis  (D = disclosed by Oracle, E = analyst estimate)'])
ROW={}; r=8
BLOCK={'Base':0,'Bull':0,'Bear':0}
for key,label,single,fmt,b,u,e,src in D:
    if single is None:
        ws.cell(row=r,column=1,value=label).font=SUB
        for cc in range(1,11): ws.cell(row=r,column=cc).fill=GFILL
        r+=1; continue
    ROW[key]=r
    ws.cell(row=r,column=1,value=label).font=BLACK
    ws.cell(row=r,column=10,value=src).font=NOTE
    ws.cell(row=r,column=10).alignment=Alignment(wrap_text=True,vertical='top')
    r+=1
n_rows=r-8
BASE_R0, BULL_R0, BEAR_R0 = 200, 320, 440
def write_block(r0,vals_idx,label,colr):
    ws.cell(row=r0-2,column=1,value=f'{label} SCENARIO INPUT BLOCK — edit here').font=SUB
    hrow(ws,r0-1,['Driver','single','FY2027','FY2028','FY2029','FY2030','FY2031'],fill=False)
    rr=r0
    for key,lab,single,fmt,b,u,e,src in D:
        if single is None: rr+=1; continue
        v=(b,u,e)[vals_idx]
        ws.cell(row=rr,column=1,value=lab).font=NOTE
        if single:
            c=ws.cell(row=rr,column=2,value=v); c.font=BLUE; c.number_format=fmt
        else:
            for j in range(5):
                c=ws.cell(row=rr,column=3+j,value=(v[j] if isinstance(v,(list,tuple)) else v))
                c.font=BLUE; c.number_format=fmt
        rr+=1
write_block(BASE_R0,0,'BASE','')
write_block(BULL_R0,1,'BULL','')
write_block(BEAR_R0,2,'BEAR','')
# ACTIVE = CHOOSE(scenario, base, bull, bear)
for key,lab,single,fmt,b,u,e,src in D:
    if single is None: continue
    r0=ROW[key]; off=r0-8
    if single:
        c=ws.cell(row=r0,column=2,value=f'=CHOOSE($B$5,B{BASE_R0+off},B{BULL_R0+off},B{BEAR_R0+off})')
        c.font=BLACK; c.number_format=fmt
    else:
        for j in range(5):
            L=CL(3+j)
            c=ws.cell(row=r0,column=3+j,value=f'=CHOOSE($B$5,{L}{BASE_R0+off},{L}{BULL_R0+off},{L}{BEAR_R0+off})')
            c.font=BLACK; c.number_format=fmt
for k in ('rpo_haircut','iaas_ndc','turn_full','util','sm_pct','wacc','new_bookings'):
    for cc in range(2,8): ws.cell(row=ROW[k],column=cc).fill=YFILL
ws.column_dimensions['A'].width=48; ws.column_dimensions['B'].width=13
for cc in 'CDEFG': ws.column_dimensions[cc].width=12
ws.column_dimensions['J'].width=95
json.dump({'ROW':ROW},open('xlsx_rows.json','w'))
wb.save('ORCL_model.xlsx'); print("part1 saved; input rows:",len(ROW))
