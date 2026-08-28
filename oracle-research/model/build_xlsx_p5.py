import openpyxl, json
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import column_index_from_string as CIX
wb=openpyxl.load_workbook('ORCL_model.xlsx')
IR=json.load(open('xlsx_rows.json'))['ROW']; P2=json.load(open('xlsx_rows2.json'))
P3=json.load(open('xlsx_rows3.json')); P4=json.load(open('xlsx_rows4.json'))
BLUE=Font(name='Arial',size=10,color='0000FF'); BLACK=Font(name='Arial',size=10)
GREEN=Font(name='Arial',size=10,color='008000'); HDR=Font(name='Arial',size=10,bold=True)
TITLE=Font(name='Arial',size=13,bold=True); SUB=Font(name='Arial',size=10,bold=True,color='444444')
NOTE=Font(name='Arial',size=8,italic=True,color='666666'); WHDR=Font(name='Arial',size=10,bold=True,color='FFFFFF')
YFILL=PatternFill('solid',fgColor='FFFF00'); GFILL=PatternFill('solid',fgColor='E8E8E8'); HFILL=PatternFill('solid',fgColor='1F3864')
CUR='$#,##0;($#,##0);-'; CUR2='$#,##0.00;($#,##0.00);-'; PCT='0.0%;(0.0%);-'; NUM='#,##0;(#,##0);-'; MULT='0.0x'
CO=['C','D','E','F','G']
def sh(n):
    ws=wb.create_sheet(n); ws.sheet_view.showGridLines=False
    ws.column_dimensions['A'].width=52
    for c in 'BCDEFGH': ws.column_dimensions[c].width=13
    ws.column_dimensions['J'].width=78
    return ws
def T(ws,t,s=None):
    ws['A1']=t; ws['A1'].font=TITLE
    if s: ws['A2']=s; ws['A2'].font=NOTE
def H(ws,r,lab):
    for i,l in enumerate(lab):
        c=ws.cell(row=r,column=1+i,value=l); c.font=WHDR; c.fill=HFILL
        c.alignment=Alignment(horizontal='center' if i else 'left')
def sec(ws,r,l):
    ws.cell(row=r,column=1,value=l).font=SUB
    for cc in range(1,9): ws.cell(row=r,column=cc).fill=GFILL
def line(ws,r,label,f,fmt=CUR,font=BLACK,note=None,bold=False,fill=None):
    c=ws.cell(row=r,column=1,value=label); c.font=HDR if bold else BLACK
    for i,col in enumerate(CO):
        cc=ws.cell(row=r,column=CIX(col),value=(f[i] if isinstance(f,(list,tuple)) else f))
        cc.font=font; cc.number_format=fmt
        if fill: cc.fill=fill
    if note:
        n=ws.cell(row=r,column=10,value=note); n.font=NOTE; n.alignment=Alignment(wrap_text=True,vertical='top')
def IN(k,c=None): return f"Inputs!$B${IR[k]}" if c is None else f"Inputs!{c}${IR[k]}"
HD=['$ millions','','FY2027E','FY2028E','FY2029E','FY2030E','FY2031E','','','Note']
IS=lambda row,c: f"Income_Statement!{c}{row}"; CD=lambda row,c: f"Capex_Depreciation!{c}{row}"
DS=lambda row,c: f"Debt_Shares!{c}{row}"

ws=sh('Cash_Flow'); T(ws,'CASH FLOW — and the three different "capex" numbers',
 'The cash statement, not the income statement, is the headline output. FY2026A: OCF $31,977m, capex $(55,663)m, FCF $(23,686)m — Oracle\'s own published figure.')
H(ws,4,HD); r=6
sec(ws,r,'OPERATING'); r+=1
line(ws,r,'Net income',['='+IS(P4['I_NI'],c) for c in CO],CUR,GREEN); F_NI=r; r+=1
line(ws,r,'+ Depreciation',['='+CD(P3['R_DEP'],c) for c in CO],CUR,GREEN); F_DEP=r; r+=1
line(ws,r,'+ Amortisation of intangible assets',['='+IS(P4['I_AM'],c) for c in CO],CUR,GREEN); F_AM=r; r+=1
line(ws,r,'+ Stock-based compensation',[f'=NonGAAP_Recon!{c}{P4["N_SBC"]}' for c in CO],CUR,GREEN); F_SBC=r; r+=1
line(ws,r,'+ Asset impairment (non-cash)',[f'={IN("impair",c)}' for c in CO],CUR,GREEN); F_IMP=r; r+=1
line(ws,r,'- Increase in net working capital',
     [f'=-({IS(P4["I_REV"],"C")}-Historicals!D13)*{IN("nwc_pct")}']+
     [f'=-({IS(P4["I_REV"],CO[i])}-{IS(P4["I_REV"],CO[i-1])})*{IN("nwc_pct")}' for i in range(1,5)]); F_NWC=r; r+=1
line(ws,r,'+ Customer prepayments',[f'={IN("cust_pre",c)}' for c in CO],CUR,GREEN,
     'FY2026A $4,592m — a brand-new line in the FY26 cash flow statement (10-K R8). These are prepayments with a significant '
     'financing component on AI contracts, i.e. customers funding Oracle\'s buildout.'); F_CP=r; r+=1
line(ws,r,'OPERATING CASH FLOW',[f'=SUM({c}{F_NI}:{c}{F_CP})' for c in CO],CUR,BLACK,'FY2026A $31,977m, +54%.',bold=True); F_OCF=r; r+=2
sec(ws,r,'THE THREE CAPEX DEFINITIONS — pin this down before quoting any cash-flow gap'); r+=1
line(ws,r,'(1) Capital expenditures, as reported',['=-'+CD(P3['R_CX'],c) for c in CO],CUR,GREEN,
     'The GAAP cash-flow-statement line. FY2026A $(55,663)m.'); F_CX=r; r+=1
line(ws,r,'      add back: short-term financing related to capex',[f'={IN("st_fin",c)}' for c in CO],CUR,GREEN,
     'FY2026A $3,345m, reported inside FINANCING activities (10-K R8).'); F_STF=r; r+=1
line(ws,r,'      add back: customer prepayments for capex',[f'={IN("cust_pre",c)}' for c in CO],CUR,GREEN,
     'FY2026A $4,592m, reported inside OPERATING activities.'); F_CPX=r; r+=1
line(ws,r,'(2) Net cash outlay for capex (Oracle non-GAAP)',[f'={c}{F_CX}+{c}{F_STF}+{c}{F_CPX}' for c in CO],CUR,BLACK,
     'ORACLE\'S OWN DEFINITION, press release Ex-99.1. FY2026A $(47,726)m. NOTE: consensus shows FY2027E free cash flow of '
     '$(47.73)bn — numerically IDENTICAL to Oracle\'s FY2026 ACTUAL net capex outlay. Treat that coincidence as a data-quality warning.',fill=YFILL)
F_NCX=r; r+=1
line(ws,r,'(3) Economic capex incl. change in unpaid capex',[f'={c}{F_CX}' for c in CO],CUR,BLACK,
     'FY2026A: $(55,663)m cash + $5,279m unpaid at year end less $2,970m at prior year end = $(57,972)m economic. Not forecast here.')
r+=1
line(ws,r,'FREE CASH FLOW (OCF less reported capex)',[f'={c}{F_OCF}+{c}{F_CX}' for c in CO],CUR,BLACK,
     'FY2026A $(23,686)m — this is the number Oracle itself publishes.',bold=True); F_FCF=r; r+=1
line(ws,r,'FREE CASH FLOW (on Oracle\'s net-capex definition)',[f'={c}{F_OCF}-{c}{F_CPX}+{c}{F_NCX}' for c in CO],CUR,BLACK,None,bold=True); F_FCF2=r; r+=1
line(ws,r,'  FCF margin',[f'={c}{F_FCF}/{IS(P4["I_REV"],c)}' for c in CO],PCT); r+=2
sec(ws,r,'FINANCING AND THE FUNDING PLUG'); r+=1
line(ws,r,'Common dividends',[f'=-{IN("div_ps",c)}*'+DS(P3['S_BAS'],c) for c in CO],CUR,BLACK,
     'DISCLOSED: $0.50/qtr declared 6/10/26 — FLAT versus FY26. After 12 consecutive years of growth, dividend growth has stopped.'); F_DIV=r; r+=1
line(ws,r,'Preferred dividends',['=-'+DS(P3['S_PD'],c) for c in CO],CUR,GREEN); F_PDIV=r; r+=1
line(ws,r,'Scheduled debt repayments',['='+DS(P3['D_MAT'],c) for c in CO],CUR,GREEN); F_MAT=r; r+=1
line(ws,r,'Equity issuance (ATM)',[f'={IN("eq_raise",c)}' for c in CO],CUR,GREEN); F_EQ=r; r+=1
line(ws,r,'Short-term financing related to capex',[f'={IN("st_fin",c)}' for c in CO],CUR,GREEN); F_STF2=r; r+=1
line(ws,r,'Target cash cushion build',[3000]*5,CUR,BLUE); F_CUSH=r; r+=1
line(ws,r,'FUNDING NEED (= new debt required)',
     [f'=-({c}{F_OCF}+{c}{F_CX}+{c}{F_DIV}+{c}{F_PDIV}+{c}{F_MAT}+{c}{F_EQ}+{c}{F_STF2})+{c}{F_CUSH}' for c in CO],CUR,BLACK,
     'Drives the new-debt plug on Debt_Shares.',bold=True); F_FUND=r; r+=1
line(ws,r,'New debt issued',['='+DS(P3['D_ISS'],c) for c in CO],CUR,GREEN); F_NEW=r; r+=1
line(ws,r,'Cash, opening',[f'={IN("cash_open")}']+[f'={CO[i-1]}{r+1}' for i in range(1,5)],CUR,BLACK); F_C0=r; r+=1
line(ws,r,'CASH, CLOSING',
     [f'={c}{F_C0}+{c}{F_OCF}+{c}{F_CX}+{c}{F_DIV}+{c}{F_PDIV}+{c}{F_MAT}+{c}{F_EQ}+{c}{F_STF2}+{c}{F_NEW}' for c in CO],CUR,BLACK,None,bold=True)
F_CASH=r; r+=2
sec(ws,r,'CREDIT METRICS'); r+=1
line(ws,r,'Net debt',['='+DS(P3['D_ND'],c) for c in CO],CUR,GREEN); r+=1
line(ws,r,'Net debt / EBITDA',[f'='+DS(P3['D_ND'],c)+'/'+IS(P4['I_EBITDA'],c) for c in CO],MULT,BLACK,
     'S&P cut Oracle to BBB- on 7/9/2026 (UNVERIFIED SECONDARY — from the user brief, not confirmed at source).'); r+=1
line(ws,r,'EBITDA / net interest expense',
     [f'=IFERROR('+IS(P4['I_EBITDA'],c)+'/('+DS(P3['D_INT'],c)+'-'+DS(P3['D_II'],c)+'),"n/m")' for c in CO],MULT,BLACK,
     'The revolving credit agreement carries a consolidated EBITDA / consolidated net interest expense covenant of 3.0x (10-K R58).',fill=YFILL)
# patch Debt_Shares placeholders
d=wb['Debt_Shares']
for c in CO:
    d[f'{c}{P3["D_ISS"]}']=f'=MAX(0,Cash_Flow!{c}{F_FUND})'
    d[f'{c}{P3["D_CASH"]}']=f'=Cash_Flow!{c}{F_CASH}'
wb.save('ORCL_model.xlsx')
json.dump(dict(F_OCF=F_OCF,F_CX=F_CX,F_FCF=F_FCF,F_FCF2=F_FCF2,F_NCX=F_NCX,F_FUND=F_FUND,F_CASH=F_CASH,
               F_NWC=F_NWC,F_CP=F_CP,F_NEW=F_NEW,F_MAT=F_MAT,F_STF2=F_STF2),open('xlsx_rows5.json','w'))
print('part5 saved: OCF',F_OCF,'FCF',F_FCF,'fund',F_FUND,'cash',F_CASH)
