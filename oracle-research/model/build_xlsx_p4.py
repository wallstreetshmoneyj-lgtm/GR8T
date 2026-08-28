import openpyxl, json
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import column_index_from_string as CIX
wb=openpyxl.load_workbook('ORCL_model.xlsx')
IR=json.load(open('xlsx_rows.json'))['ROW']; P2=json.load(open('xlsx_rows2.json')); P3=json.load(open('xlsx_rows3.json'))
BLUE=Font(name='Arial',size=10,color='0000FF'); BLACK=Font(name='Arial',size=10)
GREEN=Font(name='Arial',size=10,color='008000'); HDR=Font(name='Arial',size=10,bold=True)
TITLE=Font(name='Arial',size=13,bold=True); SUB=Font(name='Arial',size=10,bold=True,color='444444')
NOTE=Font(name='Arial',size=8,italic=True,color='666666'); WHDR=Font(name='Arial',size=10,bold=True,color='FFFFFF')
YFILL=PatternFill('solid',fgColor='FFFF00'); GFILL=PatternFill('solid',fgColor='E8E8E8'); HFILL=PatternFill('solid',fgColor='1F3864')
CUR='$#,##0;($#,##0);-'; CUR2='$#,##0.00;($#,##0.00);-'; PCT='0.0%;(0.0%);-'; NUM='#,##0;(#,##0);-'; MULT='0.0x'
CO=['C','D','E','F','G']
def sh(n):
    ws=wb.create_sheet(n); ws.sheet_view.showGridLines=False
    ws.column_dimensions['A'].width=52
    for c in 'BCDEFGH': ws.column_dimensions[c].width=13
    ws.column_dimensions['J'].width=78
    return ws
def T(ws,t,s=None):
    ws['A1']=t; ws['A1'].font=TITLE
    if s: ws['A2']=s; ws['A2'].font=NOTE
def H(ws,r,lab):
    for i,l in enumerate(lab):
        c=ws.cell(row=r,column=1+i,value=l); c.font=WHDR; c.fill=HFILL
        c.alignment=Alignment(horizontal='center' if i else 'left')
def sec(ws,r,l):
    ws.cell(row=r,column=1,value=l).font=SUB
    for cc in range(1,9): ws.cell(row=r,column=cc).fill=GFILL
def line(ws,r,label,f,fmt=CUR,font=BLACK,note=None,bold=False,fill=None):
    c=ws.cell(row=r,column=1,value=label); c.font=HDR if bold else BLACK
    for i,col in enumerate(CO):
        cc=ws.cell(row=r,column=CIX(col),value=(f[i] if isinstance(f,(list,tuple)) else f))
        cc.font=font; cc.number_format=fmt
        if fill: cc.fill=fill
    if note:
        n=ws.cell(row=r,column=10,value=note); n.font=NOTE; n.alignment=Alignment(wrap_text=True,vertical='top')
def IN(k,c=None): return f"Inputs!$B${IR[k]}" if c is None else f"Inputs!{c}${IR[k]}"
HD=['$ millions','','FY2027E','FY2028E','FY2029E','FY2030E','FY2031E','','','Note']
RS=lambda row,c: f"RPO_Schedule!{c}{row}"
CD=lambda row,c: f"Capex_Depreciation!{c}{row}"
DS=lambda row,c: f"Debt_Shares!{c}{row}"

# ==================================================== INCOME STATEMENT
ws=sh('Income_Statement'); T(ws,'INCOME STATEMENT — GAAP, FY2027E–FY2031E',
  'Depreciation sits inside cost of revenue, as Oracle reports it. That is why the buildout hits gross margin, non-GAAP EPS and GAAP EPS alike.')
H(ws,4,HD); r=6
sec(ws,r,'REVENUE'); r+=1
S=P2['SEGR']
line(ws,r,'Cloud infrastructure (IaaS)',[RS(P2['R_IAAS'],c) for c in CO],CUR,GREEN)
for c in CO: ws[f'{c}{r}']='='+RS(P2['R_IAAS'],c)
I_IA=r; r+=1
for lab,key in [('Cloud applications (SaaS)','saas_g'),('Software licence','lic_g'),('Software support','sup_g'),('Services','svc_g'),('Hardware','hw_g')]:
    line(ws,r,lab,['='+RS(S[key],c) for c in CO],CUR,GREEN)
    globals()['I_'+key]=r; r+=1
line(ws,r,'TOTAL REVENUE',[f'=SUM({c}{I_IA}:{c}{r-1})' for c in CO],CUR,BLACK,None,bold=True); I_REV=r; r+=1
line(ws,r,'  y/y growth',[f'=C{I_REV}/Historicals!D13-1']+[f'={CO[i]}{I_REV}/{CO[i-1]}{I_REV}-1' for i in range(1,5)],PCT); r+=2
sec(ws,r,'COST OF REVENUE'); r+=1
line(ws,r,'Depreciation charged to IaaS',[f'='+CD(P3['R_DEP'],c)+f'*{IN("oci_dep_sh",c)}' for c in CO],CUR,BLACK,
     'The buildout arrives here. Oracle does NOT exclude depreciation from non-GAAP.'); I_DIA=r; r+=1
line(ws,r,'IaaS cost ex-depreciation',[f'={c}{I_IA}*{IN("iaas_ndc",c)}' for c in CO],CUR,BLACK,'KEY LEVER.',fill=YFILL); I_NDC=r; r+=1
line(ws,r,'Cost of IaaS',[f'={c}{I_DIA}+{c}{I_NDC}' for c in CO]); I_CIA=r; r+=1
line(ws,r,'Cost of SaaS',[f'={c}{I_saas_g}*(1-{IN("saas_gm",c)})' for c in CO]); I_CSA=r; r+=1
line(ws,r,'Cost of software',[f'={c}{I_sup_g}*(1-{IN("sup_gm",c)})+{c}{I_lic_g}*(1-{IN("lic_gm",c)})' for c in CO]); I_CSW=r; r+=1
line(ws,r,'Cost of services',[f'={c}{I_svc_g}*(1-{IN("svc_gm",c)})' for c in CO]); I_CSV=r; r+=1
line(ws,r,'Cost of hardware',[f'={c}{I_hw_g}*(1-{IN("hw_gm",c)})' for c in CO]); I_CHW=r; r+=1
line(ws,r,'TOTAL COST OF REVENUE',[f'=SUM({c}{I_CIA}:{c}{I_CHW})' for c in CO],CUR,BLACK,None,bold=True); I_COR=r; r+=1
line(ws,r,'GROSS PROFIT',[f'={c}{I_REV}-{c}{I_COR}' for c in CO],CUR,BLACK,None,bold=True); I_GP=r; r+=1
line(ws,r,'  gross margin',[f'={c}{I_GP}/{c}{I_REV}' for c in CO],PCT,BLACK,
     'FY2026A 65.82%. Consensus FY2027E 59.98%.',bold=True); I_GM=r; r+=1
line(ws,r,'  IaaS gross margin (memo)',[f'=1-{c}{I_CIA}/{c}{I_IA}' for c in CO],PCT,BLACK,'FY2026A ~30% (derived on the Historicals tab).'); I_IGM=r; r+=2
sec(ws,r,'OPERATING EXPENSE'); r+=1
line(ws,r,'Sales & marketing',[f'={c}{I_REV}*{IN("sm_pct",c)}' for c in CO],CUR,BLACK,'FY2026A $8,331m = 12.4% of revenue.',fill=YFILL); I_SM=r; r+=1
line(ws,r,'Research & development',[f'=Historicals!D18*(1+{IN("rd_g","C")})']+[f'={CO[i-1]}{r}*(1+{IN("rd_g",CO[i])})' for i in range(1,5)]); I_RD=r; r+=1
line(ws,r,'General & administrative',[f'=Historicals!D19*(1+{IN("ga_g","C")})']+[f'={CO[i-1]}{r}*(1+{IN("ga_g",CO[i])})' for i in range(1,5)]); I_GA=r; r+=1
line(ws,r,'Amortisation of intangible assets',[f'={IN("amort",c)}' for c in CO],CUR,GREEN,
     'DISCLOSED in full: 10-K R53 gives FY27–FY31 amortisation. No estimate required.'); I_AM=r; r+=1
line(ws,r,'Restructuring, other and impairment',[f'={IN("restr",c)}+{IN("impair",c)}' for c in CO],CUR,GREEN); I_RE=r; r+=1
line(ws,r,'GAAP OPERATING INCOME',[f'={c}{I_GP}-{c}{I_SM}-{c}{I_RD}-{c}{I_GA}-{c}{I_AM}-{c}{I_RE}' for c in CO],CUR,BLACK,None,bold=True); I_OI=r; r+=1
line(ws,r,'  operating margin',[f'={c}{I_OI}/{c}{I_REV}' for c in CO],PCT); r+=2
sec(ws,r,'BELOW THE LINE'); r+=1
line(ws,r,'Interest expense',[f'=-'+DS(P3['D_INT'],c) for c in CO],CUR,GREEN,'FY2026A $(4,599)m.'); I_INT=r; r+=1
line(ws,r,'Interest income',['='+DS(P3['D_II'],c) for c in CO],CUR,GREEN); I_II=r; r+=1
line(ws,r,'Other non-operating (FX, NCI)',[-250]*5,CUR,BLUE,
     'FY2026A non-operating income was +$3,547m, but $2,811m of that was one-time investment gains (Ampere, Bloom warrants). '
     'No repeat is assumed.'); I_OT=r; r+=1
line(ws,r,'PRETAX INCOME',[f'={c}{I_OI}+{c}{I_INT}+{c}{I_II}+{c}{I_OT}' for c in CO],CUR,BLACK,None,bold=True); I_PT=r; r+=1
line(ws,r,'Provision for income taxes',[f'=-{c}{I_PT}*{IN("tax_gaap",c)}' for c in CO],CUR,BLACK,
     'FY2026A effective rate 12.6%, which included a $2,062m stock-compensation windfall benefit (10-K R80). '
     'At $151.94 versus a $345.72 52-week high, that benefit does not repeat.',fill=YFILL); I_TX=r; r+=1
line(ws,r,'NET INCOME',[f'={c}{I_PT}+{c}{I_TX}' for c in CO],CUR,BLACK,None,bold=True); I_NI=r; r+=1
line(ws,r,'Preferred dividends',['=-'+DS(P3['S_PD'],c) for c in CO],CUR,GREEN); I_PD=r; r+=1
line(ws,r,'NET INCOME AVAILABLE TO COMMON',[f'={c}{I_NI}+{c}{I_PD}' for c in CO],CUR,BLACK,None,bold=True); I_NIC=r; r+=1
line(ws,r,'Diluted shares (m)',['='+DS(P3['S_CL'],c) for c in CO],NUM,GREEN); I_SH=r; r+=1
line(ws,r,'GAAP DILUTED EPS',[f'={c}{I_NIC}/{c}{I_SH}' for c in CO],CUR2,BLACK,'FY2026A $5.83.',bold=True); I_EPS=r; r+=1
line(ws,r,'EBITDA (GAAP op income + D&A)',[f'={c}{I_OI}+'+CD(P3['R_DEP'],c)+f'+{c}{I_AM}' for c in CO],CUR,BLACK,'FY2026A $29,900m.'); I_EBITDA=r

# ==================================================== NON-GAAP RECON
ws=sh('NonGAAP_Recon'); T(ws,'GAAP <-> NON-GAAP RECONCILIATION',
  'Oracle excludes SBC, intangible amortisation and restructuring — and RE-TAXES the result at a HIGHER rate than its GAAP effective rate. That tax wedge is a real part of why non-GAAP EPS growth lags revenue growth.')
H(ws,4,['$ millions','FY2026A','FY2027E','FY2028E','FY2029E','FY2030E','FY2031E','','','Note']); r=6
sec(ws,r,'FY2026 ACTUAL BRIDGE (press release Ex-99.1, verbatim)'); r+=1
for lab,g,adj,ng,note in [
    ('Total operating expenses',46751,-8320,38431,'Add-backs: SBC $4,811m + intangible amortisation $1,671m + restructuring $1,838m = $8,320m.'),
    ('Operating income',20606,8320,28926,'GAAP margin 31%, non-GAAP margin 43%.'),
    ('Provision for income taxes',2467,3070,5537,'GAAP effective rate 12.6%; non-GAAP effective rate 19.9%. The non-GAAP rate is HIGHER.'),
    ('Net income',17087,5250,22337,''),
    ('Net income available to common',16984,5250,22234,''),
]:
    ws.cell(row=r,column=1,value=lab).font=BLACK
    for j,v in enumerate((g,adj,ng)):
        c=ws.cell(row=r,column=2+j,value=v); c.font=BLUE; c.number_format=CUR
    if note: n=ws.cell(row=r,column=10,value=note); n.font=NOTE; n.alignment=Alignment(wrap_text=True,vertical='top')
    r+=1
ws.cell(row=r-6,column=2,value=None)
for j,l in enumerate(['GAAP','Adjustment','Non-GAAP']): ws.cell(row=5,column=2+j,value=l).font=HDR
ws.cell(row=r,column=1,value='Diluted EPS').font=HDR
for j,v in enumerate((5.83,1.80,7.63)):
    c=ws.cell(row=r,column=2+j,value=v); c.font=BLUE; c.number_format=CUR2
ws.cell(row=r,column=10,value='Diluted share count is IDENTICAL on both bases: 2,914m.').font=NOTE; r+=1
ws.cell(row=r,column=1,value='Non-GAAP EPS excluding one-time investment gains').font=HDR
c=ws.cell(row=r,column=4,value=6.83); c.font=BLUE; c.number_format=CUR2; c.fill=YFILL
ws.cell(row=r,column=10,value='DISCLOSED: press release footnote 1 — the $7.63 headline includes gains on the Ampere sale and Bloom Energy warrants. '
        '$6.83 is the clean base, and it is the base management itself uses for the "+18%" FY27 growth claim.').font=NOTE
ws.cell(row=r,column=10).alignment=Alignment(wrap_text=True,vertical='top'); r+=2
sec(ws,r,'FORECAST BRIDGE'); r+=1
line(ws,r,'GAAP operating income',['=Income_Statement!'+f'{c}{I_OI}' for c in CO],CUR,GREEN); N_OI=r; r+=1
line(ws,r,'+ Stock-based compensation',[f'=Historicals!D34*(1+{IN("sbc_g","C")})']+[f'={CO[i-1]}{r}*(1+{IN("sbc_g",CO[i])})' for i in range(1,5)],CUR,BLACK,
     'FY2026A $4,811m.'); N_SBC=r; r+=1
line(ws,r,'+ Amortisation of intangible assets',[f'=Income_Statement!{c}{I_AM}' for c in CO],CUR,GREEN); N_AM=r; r+=1
line(ws,r,'+ Restructuring, other and impairment',[f'=Income_Statement!{c}{I_RE}' for c in CO],CUR,GREEN,
     'In the bear case the asset write-downs land HERE — inside the very line Oracle excludes. Non-GAAP EPS barely moves while GAAP EPS collapses.'); N_RE=r; r+=1
line(ws,r,'NON-GAAP OPERATING INCOME',[f'={c}{N_OI}+{c}{N_SBC}+{c}{N_AM}+{c}{N_RE}' for c in CO],CUR,BLACK,None,bold=True); N_NOI=r; r+=1
line(ws,r,'  non-GAAP operating margin',[f'={c}{N_NOI}/Income_Statement!{c}{I_REV}' for c in CO],PCT); r+=1
line(ws,r,'Non-GAAP pretax income',[f'=Income_Statement!{c}{I_PT}+{c}{N_SBC}+{c}{N_AM}+{c}{N_RE}' for c in CO]); N_PT=r; r+=1
line(ws,r,'Non-GAAP tax',[f'=-{c}{N_PT}*{IN("tax_ng",c)}' for c in CO],CUR,BLACK,'FY2026A 19.9%, ABOVE the 12.6% GAAP rate.'); N_TX=r; r+=1
line(ws,r,'Non-GAAP net income',[f'={c}{N_PT}+{c}{N_TX}' for c in CO]); N_NI=r; r+=1
line(ws,r,'Less preferred dividends',[f'=Income_Statement!{c}{I_PD}' for c in CO],CUR,GREEN); N_PD=r; r+=1
line(ws,r,'NON-GAAP EPS',[f'=({c}{N_NI}+{c}{N_PD})/Income_Statement!{c}{I_SH}' for c in CO],CUR2,BLACK,None,bold=True); N_EPS=r; r+=1
line(ws,r,'MEMO: FY2027 non-GAAP EPS GUIDANCE',[8.05,'','','',''],CUR2,BLUE,
     'DISCLOSED: press release 6/10/2026 — "raise our non-GAAP EPS guidance to $8.05". Guidance, not consensus.'); r+=1
line(ws,r,'MEMO: FY2027 consensus non-GAAP EPS',[8.05,'','','',''],CUR2,BLUE,
     'stockanalysis.com, 42 analysts: low $7.72 / avg $8.05 / high $8.57. Consensus sits ON guidance — that is herding, not conviction.'); r+=1
line(ws,r,'GAAP EPS (memo)',[f'=Income_Statement!{c}{I_EPS}' for c in CO],CUR2,GREEN)

wb.save('ORCL_model.xlsx')
json.dump(dict(I_REV=I_REV,I_OI=I_OI,I_NI=I_NI,I_NIC=I_NIC,I_SH=I_SH,I_EPS=I_EPS,I_EBITDA=I_EBITDA,I_AM=I_AM,
               I_PT=I_PT,I_GM=I_GM,I_IGM=I_IGM,I_IA=I_IA,I_TX=I_TX,N_SBC=N_SBC,N_EPS=N_EPS,N_NOI=N_NOI),open('xlsx_rows4.json','w'))
print('part4 saved: IS rev',I_REV,'oi',I_OI,'eps',I_EPS,'| nonGAAP eps',N_EPS)
